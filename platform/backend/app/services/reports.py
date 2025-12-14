from datetime import datetime
from pathlib import Path
from typing import Dict, Iterable, List, Optional

from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from sqlmodel import Session, select

from app.models.core import Asset, Observation, Report, SecurityEvent


def _collect_assets(session: Session, asset_ids: Optional[List[int]] = None) -> List[Asset]:
    query = select(Asset)
    if asset_ids:
        query = query.where(Asset.id.in_(asset_ids))
    return list(session.exec(query))


def _collect_events(session: Session, asset_ids: Optional[List[int]] = None) -> List[SecurityEvent]:
    query = select(SecurityEvent)
    if asset_ids:
        query = query.where(SecurityEvent.asset_id.in_(asset_ids))
    return list(session.exec(query))


def _collect_observations(session: Session, asset_ids: Optional[List[int]] = None) -> List[Observation]:
    query = select(Observation)
    if asset_ids:
        query = query.where(Observation.asset_id.in_(asset_ids))
    return list(session.exec(query))


def _render_html(title: str, assets: Iterable[Asset], events: Iterable[SecurityEvent], observations: Iterable[Observation], parameters: Dict) -> str:
    asset_rows = "".join(
        f"<tr><td>{a.id}</td><td>{a.ip}</td><td>{a.device_type or ''}</td><td>{', '.join(a.protocols)}</td></tr>"
        for a in assets
    )
    event_rows = "".join(
        f"<tr><td>{e.id}</td><td>{e.asset_id or ''}</td><td>{e.severity}</td><td>{e.attack_stage or 'n/a'}</td><td>{e.risk_score}</td><td>{e.description}</td></tr>"
        for e in events
    )
    observation_rows = "".join(
        f"<tr><td>{o.id}</td><td>{o.asset_id or ''}</td><td>{o.protocol}</td><td>{o.timestamp.isoformat()}</td></tr>"
        for o in observations
    )
    return f"""
<html>
<head><title>{title}</title></head>
<body>
<h1>{title}</h1>
<p>Generated at: {datetime.utcnow().isoformat()}Z</p>
<h2>Generation Parameters</h2>
<pre>{parameters}</pre>
<h2>Assets</h2>
<table border='1' cellspacing='0' cellpadding='4'>
<tr><th>ID</th><th>IP</th><th>Device Type</th><th>Protocols</th></tr>
{asset_rows}
</table>
<h2>Security Events</h2>
<table border='1' cellspacing='0' cellpadding='4'>
<tr><th>ID</th><th>Asset</th><th>Severity</th><th>Stage</th><th>Risk</th><th>Description</th></tr>
{event_rows}
</table>
<h2>Observations</h2>
<table border='1' cellspacing='0' cellpadding='4'>
<tr><th>ID</th><th>Asset</th><th>Protocol</th><th>Timestamp</th></tr>
{observation_rows}
</table>
</body>
</html>
"""


def _write_pdf(path: Path, title: str, assets: Iterable[Asset], events: Iterable[SecurityEvent]) -> None:
    c = canvas.Canvas(str(path), pagesize=A4)
    width, height = A4
    y = height - 50
    c.setFont("Helvetica-Bold", 16)
    c.drawString(40, y, title)
    y -= 30
    c.setFont("Helvetica", 10)
    c.drawString(40, y, f"Generated at {datetime.utcnow().isoformat()}Z")
    y -= 30
    c.setFont("Helvetica-Bold", 12)
    c.drawString(40, y, "Assets")
    c.setFont("Helvetica", 10)
    y -= 20
    for asset in assets:
        c.drawString(50, y, f"[{asset.id}] {asset.ip} ({asset.device_type or 'unknown'}) protocols: {', '.join(asset.protocols)}")
        y -= 15
        if y < 80:
            c.showPage()
            y = height - 50
    y -= 10
    c.setFont("Helvetica-Bold", 12)
    c.drawString(40, y, "Security Events")
    y -= 20
    c.setFont("Helvetica", 10)
    for event in events:
        line = (
            f"[{event.id}] asset {event.asset_id or 'n/a'} severity={event.severity}"
            f" stage={event.attack_stage or 'n/a'} risk={event.risk_score} desc={event.description}"
        )[:150]
        c.drawString(50, y, line)
        y -= 15
        if y < 80:
            c.showPage()
            y = height - 50
    c.showPage()
    c.save()


def _write_docx(path: Path, title: str, assets: Iterable[Asset], events: Iterable[SecurityEvent], observations: Iterable[Observation]) -> None:
    doc = Document()
    doc.add_heading(title, 0)
    doc.add_paragraph(f"Generated at {datetime.utcnow().isoformat()}Z")
    doc.add_heading("Assets", level=1)
    for asset in assets:
        doc.add_paragraph(f"ID {asset.id} IP {asset.ip} type {asset.device_type or 'unknown'} protocols {', '.join(asset.protocols)}")
    doc.add_heading("Security Events", level=1)
    for event in events:
        doc.add_paragraph(
            f"[{event.severity}] stage={event.attack_stage or 'n/a'} risk={event.risk_score} asset {event.asset_id or 'n/a'}: {event.description}"
        )
    doc.add_heading("Observations", level=1)
    for obs in observations:
        doc.add_paragraph(f"{obs.protocol} @ {obs.timestamp.isoformat()} for asset {obs.asset_id or 'n/a'}")
    doc.save(path)


def generate_report(
    session: Session,
    scope: Dict,
    parameters: Dict,
    fmt: str = "html",
    title: str = "ICS Forensics Report",
    output_dir: Path = Path("reports"),
) -> Report:
    output_dir.mkdir(parents=True, exist_ok=True)
    asset_ids = scope.get("assets") if scope else None
    assets = _collect_assets(session, asset_ids)
    events = _collect_events(session, asset_ids)
    observations = _collect_observations(session, asset_ids)

    timestamp = datetime.utcnow().strftime("%Y%m%dT%H%M%SZ")
    filename = f"report_{timestamp}.{fmt}"
    path = output_dir / filename

    if fmt == "html":
        html = _render_html(title, assets, events, observations, parameters)
        path.write_text(html, encoding="utf-8")
    elif fmt == "pdf":
        _write_pdf(path, title, assets, events)
    elif fmt == "docx":
        _write_docx(path, title, assets, events, observations)
    else:
        raise ValueError("Unsupported format")

    report = Report(scope=scope or {}, parameters=parameters or {}, format=fmt, file_path=str(path), created_at=datetime.utcnow())
    session.add(report)
    session.commit()
    session.refresh(report)
    return report
